#!/usr/bin/env python3
"""
Import translated text from an Excel file back into Ren'Py translation files.

Usage:
    python import_translation_xlsx.py <xlsx_file>

The xlsx file should have the same structure as to_translate.xlsx:
  - Sheet "strings":     English | Chinese | Source | Old Key
  - Sheet "chapter1":    English | Chinese | Source | Hash | Character
  - ... etc

For each row, if the Chinese cell is non-empty and different from English,
the script updates the corresponding Ren'Py translation file.
"""

import glob
import os
import re
import sys
from openpyxl import load_workbook

def quote_for_rpy(s):
    return s.replace('\\', '\\\\').replace('"', '\\"').replace('\n', '\\n').replace('\r', '\\r').replace('\t', '\\t')

def unquote_rpy(s):
    result = []
    i = 0
    while i < len(s):
        if s[i] == '\\' and i + 1 < len(s):
            nxt = s[i + 1]
            if nxt == 'n': result.append('\n')
            elif nxt == 'r': result.append('\r')
            elif nxt == 't': result.append('\t')
            elif nxt == '"': result.append('"')
            elif nxt == '\\': result.append('\\')
            else: result.append(nxt)
            i += 2
        else:
            result.append(s[i])
            i += 1
    return ''.join(result)

def import_strings(wb):
    """Import strings sheet into strings.rpy"""
    if "strings" not in wb.sheetnames:
        print("Sheet 'strings' not found, skipping.")
        return 0
    
    ws = wb["strings"]
    updates = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 4:
            continue
        english, chinese, source, old_key = row[0], row[1], row[2], row[3]
        if not chinese or not english:
            continue
        if str(chinese).strip() == str(english).strip():
            continue
        if old_key:
            updates[str(old_key)] = str(chinese)
    
    if not updates:
        print("No string translations to import.")
        return 0
    
    print(f"Importing {len(updates)} string translations into strings.rpy...")
    
    with open('game/tl/chinese_simplified/strings.rpy', 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Replace new values for matched old keys
    pattern = re.compile(
        r'(^(\s*#(?:\s*TODO)?\s*[^\n]*)\n'
        r'\s*old\s+")(.*?)"\s*\n'
        r'(\s*new\s+")(.*?)"\s*$)',
        re.MULTILINE | re.DOTALL
    )
    
    def replacer(m):
        full = m.group(1)
        comment = m.group(2)
        old_quoted = m.group(3)
        new_prefix = m.group(4)
        new_quoted = m.group(5)
        
        old_text = unquote_rpy(old_quoted)
        if old_text in updates:
            new_text = updates[old_text]
            new_quoted = quote_for_rpy(new_text)
            # Remove TODO from comment if present
            comment = comment.replace('TODO', '').strip()
            return f'{comment}\n    old "{old_quoted}"\n    new "{new_quoted}"'
        return full
    
    new_content = pattern.sub(replacer, content)
    
    if new_content != content:
        with open('game/tl/chinese_simplified/strings.rpy', 'w', encoding='utf-8') as f:
            f.write(new_content)
        print(f"  Updated strings.rpy with {len(updates)} entries")
    else:
        print("  No changes made to strings.rpy")
    
    return len(updates)


def import_dialogue(wb, sheet_name, rpy_path):
    """Import a dialogue sheet into a Ren'Py translation file."""
    if sheet_name not in wb.sheetnames:
        return 0
    
    ws = wb[sheet_name]
    updates = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 5:
            continue
        english, chinese, source, hash_id, character = row[0], row[1], row[2], row[3], row[4]
        if not chinese or not english:
            continue
        if str(chinese).strip() == str(english).strip():
            continue
        if hash_id:
            updates[str(hash_id)] = (str(character), str(chinese))
    
    if not updates:
        return 0
    
    print(f"Importing {len(updates)} translations into {rpy_path}...")
    
    if not os.path.exists(rpy_path):
        print(f"  File not found: {rpy_path}")
        return 0
    
    with open(rpy_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Pattern to match translate blocks
    block_pattern = re.compile(
        r'(translate\s+chinese_simplified\s+)([a-zA-Z0-9_]+)(:\s*\n)'
        r'(\s*\n)'
        r'(\s*#\s+)([^\n]+)(\n)'
        r'(\s+)([^\n]+)',
        re.MULTILINE
    )
    
    def replacer(m):
        full_start = m.group(1)
        hash_id = m.group(2)
        colon = m.group(3)
        blank = m.group(4)
        comment_indent = m.group(5)
        comment_line = m.group(6)
        newline = m.group(7)
        display_indent = m.group(8)
        display_line = m.group(9)
        
        if hash_id in updates:
            character, new_text = updates[hash_id]
            new_display = f'{character} "{new_text}"'
            return f'{full_start}{hash_id}{colon}{blank}{comment_indent}{comment_line}{newline}{display_indent}{new_display}'
        
        return m.group(0)
    
    new_content = block_pattern.sub(replacer, content)
    
    if new_content != content:
        with open(rpy_path, 'w', encoding='utf-8') as f:
            f.write(new_content)
        print(f"  Updated {rpy_path}")
        return len(updates)
    else:
        print(f"  No changes made to {rpy_path}")
        return 0


def main():
    if len(sys.argv) < 2:
        print("Usage: python import_translation_xlsx.py <xlsx_file>")
        sys.exit(1)
    
    xlsx_file = sys.argv[1]
    if not os.path.exists(xlsx_file):
        print(f"File not found: {xlsx_file}")
        sys.exit(1)
    
    print(f"Loading {xlsx_file}...")
    wb = load_workbook(xlsx_file)
    
    total = 0
    
    # Import strings
    total += import_strings(wb)
    
    # Import dialogue sheets
    sheet_to_file = {
        "chapter1": "game/tl/chinese_simplified/BKchapter1.rpy",
        "chapter2": "game/tl/chinese_simplified/BKchapter2.rpy",
        "chapter3": "game/tl/chinese_simplified/BKchapter3.rpy",
        "story": "game/tl/chinese_simplified/BKstory_events.rpy",
        "city": "game/tl/chinese_simplified/BKcity_events.rpy",
        "day": "game/tl/chinese_simplified/BKday_events.rpy",
        "interact": "game/tl/chinese_simplified/BKinteractions.rpy",
        "interact_free": "game/tl/chinese_simplified/BKinteractions_free.rpy",
        "help": "game/tl/chinese_simplified/BKhelp.rpy",
        "intro": "game/tl/chinese_simplified/BKintro.rpy",
        "main": "game/tl/chinese_simplified/BKmain.rpy",
        "powers": "game/tl/chinese_simplified/BKpowers.rpy",
        "security": "game/tl/chinese_simplified/BKsecurity.rpy",
        "start": "game/tl/chinese_simplified/BKstart.rpy",
        "kite1": "game/tl/chinese_simplified/kite_jobgirl 1_riddle.rpy",
        "kite2": "game/tl/chinese_simplified/kite_jobgirl 2_beach.rpy",
    }
    
    for sheet_name, rpy_path in sheet_to_file.items():
        count = import_dialogue(wb, sheet_name, rpy_path)
        total += count
    
    print(f"\n{'='*60}")
    print(f"Total translations imported: {total}")
    print(f"{'='*60}")
    print("\nNext steps:")
    print("  1. Delete .rpyc cache files: Remove-Item game/tl/chinese_simplified/*.rpyc")
    print("  2. Launch game to verify translations")

if __name__ == '__main__':
    main()
