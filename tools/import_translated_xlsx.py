#!/usr/bin/env python3
"""
Import translated text from a Google-Translate-processed Excel file back into Ren'Py translation files.

Google Translate typically replaces the original English column (A) with Chinese text,
while leaving the translation column (B) empty. This script uses row-by-row alignment
between the original to_translate.xlsx and the translated to_translated.xlsx.

Usage:
    python import_translated_xlsx.py <translated_xlsx>
"""

import glob
import os
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).parent.parent
STRINGS_PATH = ROOT / 'game/tl/chinese_simplified/strings.rpy'
TL_DIR = ROOT / 'game/tl/chinese_simplified'
XLSX_ORIG = ROOT / 'to_translate.xlsx'

def quote_for_rpy(s):
    return s.replace('\\', '\\\\').replace('"', '\\"').replace('\n', '\\n').replace('\r', '\\r').replace('\t', '\\t')

def unquote_rpy(s):
    result = []
    i = 0
    while i < len(s):
        if s[i] == '\\' and i + 1 < len(s):
            nxt = s[i + 1]
            if nxt == 'n': result.append('\n')
            elif nxt == 'r': result.append('\r')
            elif nxt == 't': result.append('\t')
            elif nxt == '"': result.append('"')
            elif nxt == '\\': result.append('\\')
            else: result.append(nxt)
            i += 2
        else:
            result.append(s[i])
            i += 1
    return ''.join(result)


def import_strings(orig_ws, trans_ws):
    """Import strings sheet into strings.rpy using row alignment."""
    updates = {}
    skipped = 0
    
    for row_idx in range(2, orig_ws.max_row + 1):
        orig_eng = orig_ws.cell(row=row_idx, column=1).value
        orig_src = orig_ws.cell(row=row_idx, column=3).value
        orig_key = orig_ws.cell(row=row_idx, column=4).value
        
        trans_a = trans_ws.cell(row=row_idx, column=1).value
        trans_b = trans_ws.cell(row=row_idx, column=2).value
        
        if not orig_eng or not orig_key:
            continue
        
        # Determine which column has the translation
        chinese = None
        if trans_b and str(trans_b).strip() and str(trans_b).strip() != str(orig_eng).strip():
            chinese = str(trans_b).strip()
        elif trans_a and str(trans_a).strip() and str(trans_a).strip() != str(orig_eng).strip():
            chinese = str(trans_a).strip()
        
        if chinese:
            updates[str(orig_key)] = chinese
        else:
            skipped += 1
    
    if not updates:
        print("No string translations to import.")
        return 0
    
    print(f"Importing {len(updates)} string translations into strings.rpy (skipped {skipped} empty rows)...")
    
    with open(STRINGS_PATH, 'r', encoding='utf-8') as f:
        content = f.read()
    
    pattern = re.compile(
        r'(^(\s*#(?:\s*TODO)?\s*[^\n]*)\n'
        r'\s*old\s+"(.*?)(?<!\\)"\s*\n'
        r'(\s*new\s+")(.*?)(?<!\\)"\s*$)',
        re.MULTILINE | re.DOTALL
    )
    
    match_count = 0
    def replacer(m):
        nonlocal match_count
        full = m.group(1)
        comment = m.group(2)
        old_quoted = m.group(3)
        new_prefix = m.group(4)
        new_quoted = m.group(5)
        
        old_text = unquote_rpy(old_quoted)
        if old_text in updates:
            new_text = updates[old_text]
            new_quoted = quote_for_rpy(new_text)
            comment = comment.replace('TODO', '').strip()
            match_count += 1
            return f'    {comment}\n    old "{old_quoted}"\n    new "{new_quoted}"\n'
        return full
    
    new_content = pattern.sub(replacer, content)
    
    if new_content != content:
        with open(STRINGS_PATH, 'w', encoding='utf-8') as f:
            f.write(new_content)
        print(f"  Updated strings.rpy ({match_count} entries matched and updated)")
    else:
        print("  No changes made to strings.rpy")
    
    return match_count


def import_dialogue(orig_ws, trans_ws, rpy_path):
    """Import a dialogue sheet using row alignment."""
    updates = {}
    
    for row_idx in range(2, orig_ws.max_row + 1):
        orig_eng = orig_ws.cell(row=row_idx, column=1).value
        orig_hash = orig_ws.cell(row=row_idx, column=4).value
        orig_char = orig_ws.cell(row=row_idx, column=5).value
        
        trans_a = trans_ws.cell(row=row_idx, column=1).value
        trans_b = trans_ws.cell(row=row_idx, column=2).value
        
        if not orig_eng or not orig_hash:
            continue
        
        chinese = None
        if trans_b and str(trans_b).strip() and str(trans_b).strip() != str(orig_eng).strip():
            chinese = str(trans_b).strip()
        elif trans_a and str(trans_a).strip() and str(trans_a).strip() != str(orig_eng).strip():
            chinese = str(trans_a).strip()
        
        if chinese:
            character = str(orig_char).strip() if orig_char else ""
            updates[str(orig_hash).strip()] = (character, chinese)
    
    if not updates:
        return 0
    
    print(f"Importing {len(updates)} translations into {rpy_path}...")
    
    if not os.path.exists(rpy_path):
        print(f"  File not found: {rpy_path}")
        return 0
    
    with open(rpy_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    block_pattern = re.compile(
        r'(translate\s+chinese_simplified\s+)([a-zA-Z0-9_]+)(:\s*\n)'
        r'(\s*\n)'
        r'(\s*#\s+)([^\n]+)(\n)'
        r'(\s+)([^\n]+)',
        re.MULTILINE
    )
    
    match_count = 0
    def replacer(m):
        nonlocal match_count
        full_start = m.group(1)
        hash_id = m.group(2)
        colon = m.group(3)
        blank = m.group(4)
        comment_indent = m.group(5)
        comment_line = m.group(6)
        newline = m.group(7)
        display_indent = m.group(8)
        display_line = m.group(9)
        
        if hash_id in updates:
            character, new_text = updates[hash_id]
            # Try to preserve original character tag if ours is empty
            display_match = re.match(r'(\S+)\s+"(.*)"$', display_line)
            if display_match and not character:
                character = display_match.group(1)
            if not character:
                character = "narrator"
            new_display = f'{character} "{new_text}"'
            match_count += 1
            return f'{full_start}{hash_id}{colon}{blank}{comment_indent}{comment_line}{newline}{display_indent}{new_display}'
        
        return m.group(0)
    
    new_content = block_pattern.sub(replacer, content)
    
    if new_content != content:
        with open(rpy_path, 'w', encoding='utf-8') as f:
            f.write(new_content)
        print(f"  Updated {rpy_path} ({match_count} entries)")
        return match_count
    else:
        print(f"  No changes made to {rpy_path}")
        return 0


def main():
    if len(sys.argv) < 2:
        print("Usage: python import_translated_xlsx.py <translated_xlsx>")
        sys.exit(1)
    
    translated_file = sys.argv[1]
    original_file = XLSX_ORIG
    
    if not os.path.exists(translated_file):
        print(f"File not found: {translated_file}")
        sys.exit(1)
    
    if not os.path.exists(original_file):
        print(f"Original file not found: {original_file}")
        sys.exit(1)
    
    print(f"Loading original: {original_file}")
    orig_wb = load_workbook(original_file)
    
    print(f"Loading translated: {translated_file}")
    trans_wb = load_workbook(translated_file)
    
    if len(orig_wb.sheetnames) != len(trans_wb.sheetnames):
        print(f"ERROR: Sheet count mismatch! Original: {len(orig_wb.sheetnames)}, Translated: {len(trans_wb.sheetnames)}")
        sys.exit(1)
    
    total = 0
    
    # Import strings
    if len(orig_wb.sheetnames) > 0:
        total += import_strings(orig_wb[orig_wb.sheetnames[0]], trans_wb[trans_wb.sheetnames[0]])
    
    # Import dialogue sheets
    sheet_to_file = {
        "chapter1": str(ROOT / "game/tl/chinese_simplified/BKchapter1.rpy"),
        "chapter2": str(ROOT / "game/tl/chinese_simplified/BKchapter2.rpy"),
        "chapter3": str(ROOT / "game/tl/chinese_simplified/BKchapter3.rpy"),
        "story": str(ROOT / "game/tl/chinese_simplified/BKstory_events.rpy"),
        "city": str(ROOT / "game/tl/chinese_simplified/BKcity_events.rpy"),
        "day": str(ROOT / "game/tl/chinese_simplified/BKday_events.rpy"),
        "interact": str(ROOT / "game/tl/chinese_simplified/BKinteractions.rpy"),
        "interact_free": str(ROOT / "game/tl/chinese_simplified/BKinteractions_free.rpy"),
        "help": str(ROOT / "game/tl/chinese_simplified/BKhelp.rpy"),
        "intro": str(ROOT / "game/tl/chinese_simplified/BKintro.rpy"),
        "main": str(ROOT / "game/tl/chinese_simplified/BKmain.rpy"),
        "powers": str(ROOT / "game/tl/chinese_simplified/BKpowers.rpy"),
        "security": str(ROOT / "game/tl/chinese_simplified/BKsecurity.rpy"),
        "start": str(ROOT / "game/tl/chinese_simplified/BKstart.rpy"),
        "kite1": str(ROOT / "game/tl/chinese_simplified/kite_jobgirl 1_riddle.rpy"),
        "kite2": str(ROOT / "game/tl/chinese_simplified/kite_jobgirl 2_beach.rpy"),
    }
    
    for idx in range(1, len(orig_wb.sheetnames)):
        orig_name = orig_wb.sheetnames[idx]
        trans_name = trans_wb.sheetnames[idx]
        
        if orig_name in sheet_to_file:
            count = import_dialogue(orig_wb[orig_name], trans_wb[trans_name], sheet_to_file[orig_name])
            total += count
    
    print(f"\n{'='*60}")
    print(f"Total translations imported: {total}")
    print(f"{'='*60}")
    
    # Delete rpyc cache
    print("\nDeleting .rpyc cache files...")
    for rpyc in glob.glob(str(ROOT / 'game/tl/chinese_simplified/*.rpyc')):
        os.remove(rpyc)
        print(f"  Removed {rpyc}")
    
    print("\nDone! Launch the game to verify translations.")

if __name__ == '__main__':
    main()
