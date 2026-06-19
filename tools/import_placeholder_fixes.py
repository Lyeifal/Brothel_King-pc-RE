#!/usr/bin/env python3
"""
Import fixed translations from placeholder_mismatches_<language>.xlsx
back into the corresponding .rpy translation files.

Skips rows where placeholders in the fixed translation don't match the original.
"""
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Install with: pip install openpyxl")
    sys.exit(1)

ROOT = Path(r"C:\Users\akxls\Documents\Code\BK\Brothel_King-pc")
TL_DIR = ROOT / "game" / "tl" / "chinese_simplified"

PRINTF_RE = re.compile(r"%[+-]?\d*(?:\.\d+)?[sdif]")
INTERPOLATION_RE = re.compile(r"\[[A-Za-z_][A-Za-z0-9_]*(?:\![A-Za-z]+|:[^\]]+)?\]")


def extract_placeholders(text):
    if text is None:
        return []
    printf = [m.group(0) for m in PRINTF_RE.finditer(text)]
    interp = INTERPOLATION_RE.findall(text)
    return sorted(printf + interp)


def parse_quoted_string(s):
    s = s.strip()
    if s.startswith('"""') and s.endswith('"""'):
        return s[3:-3]
    if s.startswith('"') and s.endswith('"'):
        return s[1:-1].replace('\\"', '"').replace('\\n', '\n').replace('\\\\', '\\')
    return s


def format_quoted_string(text):
    if text is None:
        return '""'
    if '\n' in text or '\r' in text:
        safe = text.replace('"""', '\\"""')
        return f'"""{safe}"""'
    else:
        safe = text.replace('"', '\\"').replace('\\', '\\\\')
        return f'"{safe}"'


def main():
    language = sys.argv[1] if len(sys.argv) > 1 else "chinese_simplified"
    excel_file = ROOT / f"placeholder_mismatches_{language}.xlsx"

    if not excel_file.exists():
        print(f"Excel file not found: {excel_file}")
        return

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active

    fixes_by_file = {}
    skipped = 0
    placeholder_mismatch = 0

    for row in range(2, ws.max_row + 1):
        file_cell = ws.cell(row, 2).value
        english = ws.cell(row, 4).value
        chinese = ws.cell(row, 5).value

        if not english or not chinese:
            skipped += 1
            continue

        old_ph = extract_placeholders(english)
        new_ph = extract_placeholders(chinese)
        if old_ph != new_ph:
            placeholder_mismatch += 1
            print(f"Placeholder mismatch (skipping): {file_cell}:{ws.cell(row, 3).value}")
            print(f"  old: {english!r}")
            print(f"  new: {chinese!r}")
            print(f"  expected: {old_ph}, actual: {new_ph}")
            continue

        # Normalize file path
        if file_cell:
            file_cell = str(file_cell).replace('\\', '/')
            prefix = 'game/tl/chinese_simplified/'
            if file_cell.startswith(prefix):
                file_cell = file_cell[len(prefix):]

        fixes_by_file.setdefault(file_cell, []).append((english, chinese))

    applied = 0
    not_found = 0

    block_pattern = re.compile(
        r'^(\s*#\s*[^\n]*\n)?'
        r'(\s*old\s+)("""[\s\S]*?"""|"(?:[^"\\]|\\.)*")\n'
        r'(\s*new\s+)("""[\s\S]*?"""|"(?:[^"\\]|\\.)*")\n',
        re.MULTILINE,
    )

    for file_cell, fixes in fixes_by_file.items():
        fpath = TL_DIR / file_cell
        if not fpath.exists():
            print(f"File not found: {fpath}")
            skipped += len(fixes)
            continue

        content = fpath.read_text(encoding='utf-8')
        original = content

        eng_to_ch = {english: chinese for english, chinese in fixes}

        def replace_block(match):
            old_prefix = match.group(2)
            old_val = match.group(3)
            new_prefix = match.group(4)
            new_val = match.group(5)

            old_text = parse_quoted_string(old_val)
            new_text = parse_quoted_string(new_val)

            if old_text in eng_to_ch and new_text != eng_to_ch[old_text]:
                return f"{old_prefix}{old_val}\n{new_prefix}{format_quoted_string(eng_to_ch[old_text])}\n"
            return match.group(0)

        new_content = block_pattern.sub(replace_block, content)

        if new_content != original:
            fpath.write_text(new_content, encoding='utf-8')

        # Count applied
        for english, chinese in eng_to_ch.items():
            formatted = format_quoted_string(chinese)
            escaped_old = re.escape(english).replace('\n', '\\n')
            if re.search(r'old\s+"' + escaped_old + r'"\s*\n\s*new\s+' + re.escape(formatted), new_content):
                applied += 1
            else:
                not_found += 1
                print(f"  Could not confirm fix: {fpath}:{english[:60]!r}")

    print(f"\nApplied: {applied}")
    print(f"Placeholder mismatches skipped: {placeholder_mismatch}")
    print(f"Blocks not found/unconfirmed: {not_found}")
    print(f"Rows skipped (empty): {skipped}")


if __name__ == "__main__":
    main()
