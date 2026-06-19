import re
import os
from pathlib import Path
import openpyxl

# Project root is parent of tools directory
ROOT = Path(__file__).parent.parent.parent.parent
STRINGS_PATH = ROOT / 'game/tl/chinese_simplified/strings.rpy'
TL_DIR = ROOT / 'game/tl/chinese_simplified'


# Strings that genuinely don't need translation (symbols, keys, codes, dynamic vars)
NO_TRANSLATE_NEEDED = {
    '#A6DEEE',
    '...',
    '69.',
    '<',
    '>',
    'B', 'C', 'L', 'S', 'X',
    'Ctrl',
    'XXX',
    'XXX{#1}',
    'd3', 'd3 + -1', 'd3 + -2', 'd3 + 1', 'd3 + 3',
    'd4 + -1',
    'd5',
    'd6', 'd6 + 2',
    'dom',
    'es',
    'g: [MC.name]...',
    '{#auto_page}A',
    '{#file_time}%A, %B %d %Y, %H:%M',
    '{#file_time}%A, %B %d, %H:%M',
    '{#quick_page}Q',
    '{0} {1}',
    '{b}JP{/b}',
    '{b}XP{/b}',
    '{b}Xp{/b}',
    '{color=[c_cream]}',
    '{color=[c_lightgreen]}',
    '{color=[c_orange]}',
    '{size=-1}',
    '\u25b2{image=img_gold}',
    '\u4e2d\u6587',
    '\u7b80\u4f53\u4e2d\u6587',
    '{color=[c_red]}\u2719{/color}',
    '',
    # Dynamic variables - translated elsewhere
    '[cust.reputation_comment!t]',
    '[girl.origin].',
    '[pow.sanity_lvl]',
    # Pure format templates (no actual English text)
    '\\n%s {color=[c_green]}+%s{/color}',
    '\\n%s {color=[c_red]}%s{/color}',
    '\\n%s: %s',
    '\\n(%s ',
    '\\nJP {color=[c_orange]}+%s{/color}',
    '%s',
    '%s ',
    '%s %s',
    '%s (%s)',
    '%s ({image=img_cust} %i/%i)',
    '%s%s',
    '%s%s%s',
    '%s: %s',
    '{b}%s{/b}',
    '{b}%s{/b} - {i}%s',
    '{b}%s{/b}: ',
    '{b}%s{/b}: %s',
    '{color=#1FCB4A}%s{/color}',
    '{color=#7CEB98}%s{/color}',
    '{color=#BDF4CB}%s{/color}',
    '{color=#F70000}%s{/color}',
    '{color=#FF2626}%s{/color}',
    '{color=#FF5353}%s{/color}',
    '{color=#FF8E8E}%s{/color}',
    '{color=#FFB5B5}%s{/color}',
    '{color=[c_darkblue]}%s{/color}',
    '{color=[c_darkred]}%s{/color}',
    '{color=[c_emerald]}%s{/color}',
    '{color=[c_lightblue]}%s{/color}',
    '{color=[c_orange]}%s{/color}',
    '{color=[c_white]}%s{/color}',
}


def has_chinese(text):
    """Check if text contains CJK characters or fullwidth punctuation (indicates translation)."""
    if not text:
        return False
    for ch in text:
        if '\u4e00' <= ch <= '\u9fff':
            return True
        if '\u3040' <= ch <= '\u309f':
            return True
        if '\u30a0' <= ch <= '\u30ff':
            return True
        if '\uac00' <= ch <= '\ud7af':
            return True
        if '\u3000' <= ch <= '\u303f':
            return True
        if '\uff00' <= ch <= '\uffef':
            return True
    return False


def extract_quoted(text):
    """Extract content from first quoted string using proper escape handling."""
    match = re.search(r'"((?:\\.|[^"\\])*)"', text)
    if match:
        return match.group(1)
    return ''


def unquote_rpy(text):
    """Unescape Ren'Py string."""
    text = text.replace('\\n', '\n')
    text = text.replace('\\"', '"')
    text = text.replace('\\\\', '\\')
    return text


def find_untranslated_strings():
    with open(STRINGS_PATH, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    untranslated = []
    total = 0
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('old '):
            total += 1
            old_val = extract_quoted(line)
            if i+1 < len(lines) and lines[i+1].strip().startswith('new '):
                new_val = extract_quoted(lines[i+1].strip())
                if not new_val or new_val == old_val:
                    # Skip entries that genuinely don't need translation
                    if unquote_rpy(old_val) in NO_TRANSLATE_NEEDED or old_val in NO_TRANSLATE_NEEDED:
                        i += 1
                        continue
                    comment = ''
                    for j in range(i-1, max(-1, i-5), -1):
                        if lines[j].strip().startswith('#'):
                            comment = lines[j].strip()
                            break
                    untranslated.append({
                        'english': unquote_rpy(old_val),
                        'chinese': '',
                        'source': comment.lstrip('#').strip(),
                        'old_key': unquote_rpy(old_val),
                    })
            i += 1
        i += 1

    return untranslated, total


def parse_translate_blocks(content, source_hint):
    """Parse translate blocks from a Ren'Py translation file."""
    blocks = []
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        match = re.match(r'translate\s+chinese_simplified\s+(\w+)\s*:\s*$', stripped)
        if match:
            label = match.group(1)
            j = i + 1
            comment_line = None
            dialogue_line = None
            while j < len(lines) and j < i + 10:
                s = lines[j].strip()
                if s.startswith('# '):
                    comment_line = s[2:]
                elif s and not s.startswith('#'):
                    dialogue_line = s
                    break
                j += 1
            if comment_line and dialogue_line:
                orig_text = extract_quoted(comment_line)
                trans_text = extract_quoted(dialogue_line)
                speaker_match = re.match(r'((?:\S+\s+)*)"', dialogue_line)
                speaker = speaker_match.group(1).strip() if speaker_match else ''
                blocks.append({
                    'label': label,
                    'speaker': speaker,
                    'english': unquote_rpy(orig_text),
                    'chinese': unquote_rpy(trans_text),
                    'source': source_hint,
                })
        i += 1
    return blocks


def find_untranslated_dialogues():
    tl_dir = TL_DIR
    all_blocks = []

    file_to_sheet = {
        'BKchapter1.rpy': 'chapter1',
        'BKchapter2.rpy': 'chapter2',
        'BKchapter3.rpy': 'chapter3',
        'BKstory_events.rpy': 'story',
        'BKcity_events.rpy': 'city',
        'BKday_events.rpy': 'day',
        'BKinteractions.rpy': 'interact',
        'BKinteractions_free.rpy': 'interact_free',
        'BKhelp.rpy': 'help',
        'BKintro.rpy': 'intro',
        'BKmain.rpy': 'main',
        'BKpowers.rpy': 'powers',
        'BKsecurity.rpy': 'security',
        'BKstart.rpy': 'start',
        'kite_jobgirl 1_riddle.rpy': 'kite1',
        'kite_jobgirl 2_beach.rpy': 'kite2',
    }

    sheet_blocks = {k: [] for k in file_to_sheet.values()}

    for rpy_file, sheet_name in file_to_sheet.items():
        path = tl_dir / rpy_file
        if not path.exists():
            continue
        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()
        blocks = parse_translate_blocks(content, str(path))
        for b in blocks:
            if not has_chinese(b['chinese']) or b['chinese'] == b['english']:
                sheet_blocks[sheet_name].append(b)
        all_blocks.extend(blocks)

    return sheet_blocks, all_blocks


def generate_xlsx(string_items, sheet_blocks):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('strings')
    ws.append(['English (原文)', 'Chinese (翻译)', 'Source', 'Old Key'])
    for item in string_items:
        ws.append([item['english'], item['chinese'], item['source'], item['old_key']])

    sheet_order = ['chapter1', 'chapter2', 'chapter3', 'story', 'city', 'day',
                   'interact', 'interact_free', 'help', 'intro', 'main',
                   'powers', 'security', 'start', 'kite1', 'kite2']

    for sheet_name in sheet_order:
        ws = wb.create_sheet(sheet_name)
        ws.append(['English (原文)', 'Chinese (翻译)', 'Source', 'Speaker', 'Label'])
        for b in sheet_blocks.get(sheet_name, []):
            ws.append([b['english'], b['chinese'], b['source'], b['speaker'], b['label']])

    output_path = ROOT / 'temp' / 'translations' / 'to_translate_remaining_v2.xlsx'
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f'Saved {output_path}')


if __name__ == '__main__':
    s_untrans, s_total = find_untranslated_strings()
    print(f'strings.rpy: {len(s_untrans)}/{s_total} untranslated')

    sheet_blocks, all_blocks = find_untranslated_dialogues()
    d_untrans = sum(len(v) for v in sheet_blocks.values())
    d_total = len(all_blocks)
    print(f'Dialogue files: {d_untrans}/{d_total} untranslated')

    for name, blocks in sheet_blocks.items():
        print(f'  {name}: {len(blocks)} untranslated')

    generate_xlsx(s_untrans, sheet_blocks)

    summary_path = ROOT / 'temp' / 'translations' / 'untranslated_summary.txt'
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    with open(summary_path, 'w', encoding='utf-8') as f:
        f.write(f'strings.rpy: {len(s_untrans)}/{s_total}\n')
        f.write(f'Dialogue files: {d_untrans}/{d_total}\n')
        f.write(f'Total: {len(s_untrans) + d_untrans}\n\n')
        for item in s_untrans[:30]:
            f.write(f'string | {item["source"]} | {item["english"][:80]}\n')
        for sheet_name, blocks in sheet_blocks.items():
            for b in blocks[:5]:
                f.write(f'{sheet_name} | {b["label"]} | {b["speaker"]} | {b["english"][:80]}\n')
