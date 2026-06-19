import re
import sys
import time
from pathlib import Path
from deep_translator import GoogleTranslator
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).parent.parent.parent.parent
XLSX_PATH = ROOT / 'temp' / 'translations' / 'to_translate_remaining.xlsx'
TL_DIR = ROOT / 'game/tl/chinese_simplified'

PLACEHOLDER_FMT = '<PH_{:d}>'

PROTECT_PATTERNS = [
    r'\{[^}]+\}',
    r'\[\w+\.\w*\]',
    r'\[\w+\]',
    r'%[sdif]',
    r'\\n',
    r'\[emo_[^\]]+\]',
]

COMBINED_PATTERN = re.compile('|'.join(f'({p})' for p in PROTECT_PATTERNS))


def has_cjk_or_fullwidth(text):
    if not text:
        return False
    for ch in text:
        if '\u4e00' <= ch <= '\u9fff':
            return True
        if '\u3040' <= ch <= '\u309f':
            return True
        if '\u30a0' <= ch <= '\u30ff':
            return True
        if '\uac00' <= ch <= '\ud7af':
            return True
        if '\u3000' <= ch <= '\u303f':
            return True
        if '\uff00' <= ch <= '\uffef':
            return True
    return False


def protect_placeholders(text):
    placeholders = []
    counter = [0]
    
    def replace(m):
        ph = PLACEHOLDER_FMT.format(counter[0])
        counter[0] += 1
        placeholders.append(m.group(0))
        return ph
    
    protected = COMBINED_PATTERN.sub(replace, text)
    return protected, placeholders


def restore_placeholders(text, placeholders):
    for i, ph in enumerate(placeholders):
        text = text.replace(PLACEHOLDER_FMT.format(i), ph, 1)
    return text


def translate_text(translator, text):
    if not text or not text.strip():
        return text
    
    stripped = text.strip()
    if set(stripped).issubset({'.', ' ', '!', '?', ',', ';', ':', '-', "'", '"', '(', ')', '[', ']', '{', '}', '\\', 'n', '/'}):
        return stripped
    
    if has_cjk_or_fullwidth(stripped):
        return stripped
    
    protected, placeholders = protect_placeholders(text)
    
    try:
        translated = translator.translate(protected)
        time.sleep(0.15)
    except Exception as e:
        print(f'    Translate error ({type(e).__name__}): {str(e)[:60]}')
        return None
    
    if translated is None:
        print(f'    Translate returned None for: {text[:50]}')
        return None
    
    restored = restore_placeholders(translated, placeholders)
    return restored


def import_dialogue(translations, rpy_path):
    with open(rpy_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    match_count = 0
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        match = re.match(r'translate\s+chinese_simplified\s+(\w+)\s*:\s*$', stripped)
        if match:
            label = match.group(1)
            if label in translations:
                trans_text = translations[label]
                j = i + 1
                while j < len(lines):
                    s = lines[j].strip()
                    if s and not s.startswith('#'):
                        quoted_match = re.search(r'"((?:\\.|[^"\\])*)"', s)
                        if quoted_match:
                            quoted = quoted_match.group(1)
                            new_quoted = trans_text.replace('\\', '\\\\').replace('"', '\\"').replace('\n', '\\n')
                            lines[j] = lines[j].replace(f'"{quoted}"', f'"{new_quoted}"', 1)
                            match_count += 1
                        break
                    j += 1
        i += 1
    
    if match_count > 0:
        with open(rpy_path, 'w', encoding='utf-8') as f:
            f.writelines(lines)
        print(f'  Updated {rpy_path.name} ({match_count} entries)')
    else:
        print(f'  No matches in {rpy_path.name}')
    
    return match_count


def main():
    translator = GoogleTranslator(source='en', target='zh-CN')
    wb = openpyxl.load_workbook(XLSX_PATH)
    
    file_map = {
        'chapter1': TL_DIR / 'BKchapter1.rpy',
        'chapter2': TL_DIR / 'BKchapter2.rpy',
        'chapter3': TL_DIR / 'BKchapter3.rpy',
        'story': TL_DIR / 'BKstory_events.rpy',
        'city': TL_DIR / 'BKcity_events.rpy',
        'day': TL_DIR / 'BKday_events.rpy',
        'interact': TL_DIR / 'BKinteractions.rpy',
        'interact_free': TL_DIR / 'BKinteractions_free.rpy',
        'help': TL_DIR / 'BKhelp.rpy',
        'intro': TL_DIR / 'BKintro.rpy',
        'main': TL_DIR / 'BKmain.rpy',
        'powers': TL_DIR / 'BKpowers.rpy',
        'security': TL_DIR / 'BKsecurity.rpy',
        'start': TL_DIR / 'BKstart.rpy',
        'kite1': TL_DIR / 'kite_jobgirl 1_riddle.rpy',
        'kite2': TL_DIR / 'kite_jobgirl 2_beach.rpy',
    }
    
    total_imported = 0
    total_failed = 0
    total_skipped = 0
    
    for sheet_name, rpy_path in file_map.items():
        if sheet_name not in wb.sheetnames:
            continue
        
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))[1:]
        
        translations = {}
        sheet_failed = 0
        sheet_skipped = 0
        sheet_already = 0
        
        print(f'\n=== {sheet_name} ({len(rows)} rows) ===')
        
        for row_idx, row in enumerate(rows, 2):
            en = str(row[0]) if row and row[0] else ''
            cn = str(row[1]) if len(row) > 1 and row[1] else ''
            label = str(row[4]) if len(row) > 4 and row[4] else ''
            
            if not en.strip() or not label.strip():
                continue
            
            # Skip if already translated (CN has real translation)
            if cn.strip() and cn.strip() != en.strip() and has_cjk_or_fullwidth(cn):
                sheet_already += 1
                continue
            
            translated = translate_text(translator, en.strip())
            
            if translated is None:
                sheet_failed += 1
                total_failed += 1
                continue
            
            if translated.strip() == en.strip():
                sheet_skipped += 1
                total_skipped += 1
                continue
            
            translations[label.strip()] = translated.strip()
        
        if translations:
            print(f'  To import: {len(translations)}, Already: {sheet_already}, Failed: {sheet_failed}, Skipped: {sheet_skipped}')
            imported = import_dialogue(translations, rpy_path)
            total_imported += imported
        else:
            print(f'  Nothing to import (Already: {sheet_already}, Failed: {sheet_failed}, Skipped: {sheet_skipped})')
    
    print(f'\n{"="*60}')
    print(f'Total imported: {total_imported}')
    if total_failed > 0:
        print(f'Total failed: {total_failed}')
    if total_skipped > 0:
        print(f'Total skipped (same): {total_skipped}')
    print(f'{"="*60}')
    
    print('\nDeleting .rpyc cache...')
    for rpyc in TL_DIR.glob('*.rpyc'):
        rpyc.unlink()
    print('Done!')


if __name__ == '__main__':
    main()
