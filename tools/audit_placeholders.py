#!/usr/bin/env python3
"""
Audit placeholder integrity across all chinese_simplified translation files.

Checks that translated strings preserve the same placeholders as the original:
- printf-style: %s, %d, %i, %f, etc.
- Ren'Py interpolation: [var], [var!t], [var:...]

Reports mismatches without modifying files.
"""
import re
import sys
from pathlib import Path

# Some Windows consoles use GBK, which cannot print every Unicode character.
# Force UTF-8 output and replace unencodable characters rather than crashing.
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except AttributeError:  # Python < 3.7
    pass

ROOT = Path(__file__).resolve().parent.parent
TL_DIR = ROOT / "game" / "tl" / "chinese_simplified"

# Regex patterns
PRINTF_RE = re.compile(r"%[+-]?\d*(\.\d+)?[sdif]")
INTERPOLATION_RE = re.compile(r"\[[A-Za-z_][A-Za-z0-9_]*(?:\![A-Za-z]+|:[^\]]+)?\]")


def extract_placeholders(text):
    """Return a sorted list of placeholders in text."""
    # Strip escaped percent signs (%%) first: they render as a literal '%'
    # and must NOT count as printf placeholders. Without this, a translator
    # mistyping "%%s" instead of "%s" still scans as one real %s and the
    # missing-argument crash (TypeError at "%s" % (a, b)) goes unreported.
    printf_text = text.replace("%%", "")
    printf = [m.group(0) for m in PRINTF_RE.finditer(printf_text)]
    interp = INTERPOLATION_RE.findall(text)
    # Normalize: %s and [var]
    return sorted(printf + interp)


STRING_RE = re.compile(
    r'(?P<triple>"""[\s\S]*?""")|(?P<single>"(?:[^"\\]|\\.)*")'
)


def parse_translation_blocks(content):
    """Yield (old_text, new_text, line_number) from an .rpy translation file.

    Handles multi-line triple-quoted strings and ignores inline comments.
    """
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if stripped.startswith('old '):
            old_match = match_string_literal(lines, i)
            if old_match and old_match['end_line'] + 1 < len(lines):
                new_line = old_match['end_line'] + 1
                new_match = match_string_literal(lines, new_line)
                if new_match:
                    old_text = unquote(old_match['raw'])
                    new_text = unquote(new_match['raw'])
                    yield old_text, new_text, new_line + 1
                    i = new_match['end_line']
        i += 1


def match_string_literal(lines, start_idx):
    """Match a Ren'Py string literal starting at lines[start_idx].

    Returns a dict with 'raw' and 'end_line', or None if no string found.
    """
    line = lines[start_idx]
    # Strip leading whitespace and optional keyword (old/new)
    m = re.match(r'^\s*(?:old|new)?\s*(.*)', line)
    if not m:
        return None
    remainder = m.group(1)

    if remainder.startswith('"""'):
        # Multi-line triple-quoted string
        raw = remainder
        end_idx = start_idx
        while True:
            if raw.endswith('"""') and len(raw) > 3:
                return {'raw': raw, 'end_line': end_idx}
            end_idx += 1
            if end_idx >= len(lines):
                return None
            raw += '\n' + lines[end_idx]
    elif remainder.startswith('"'):
        # Single-line double-quoted string
        m = STRING_RE.match(remainder)
        if m and m.group('single'):
            return {'raw': m.group('single'), 'end_line': start_idx}
    return None


def unquote(raw):
    """Unquote a Ren'Py string literal (single or triple quoted)."""
    raw = raw.strip()
    if raw.startswith('"""') and raw.endswith('"""'):
        return raw[3:-3]
    if raw.startswith('"') and raw.endswith('"'):
        return raw[1:-1].replace('\\"', '"').replace('\\n', '\n').replace('\\\\', '\\')
    return raw


def export_xlsx(mismatches, language):
    try:
        import openpyxl
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    except ImportError:
        print("openpyxl not installed; skipping xlsx export.")
        return

    def clean(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "placeholder_mismatches"
    ws.append(["", "File", "Line", "English (DO NOT MODIFY)", f"{language} Translation (fix placeholders)", "Expected Placeholders", "Actual Placeholders"])
    for m in mismatches:
        ws.append(["", str(m['file']), m['line'], clean(m['old']), clean(m['new']), ", ".join(m['old_ph']), ", ".join(m['new_ph'])])
    output = ROOT / f"placeholder_mismatches_{language}.xlsx"
    wb.save(output)
    print(f"Exported to: {output}")


def main():
    mismatches = []
    checked = 0

    for rpy_file in sorted(TL_DIR.rglob("*.rpy")):
        content = rpy_file.read_text(encoding='utf-8')
        for old_text, new_text, line_no in parse_translation_blocks(content):
            checked += 1
            # Empty source string maps to empty translation by definition.
            if not old_text:
                continue
            old_ph = extract_placeholders(old_text)
            if not new_text:
                # Empty translation is dangerous when the source string is used
                # with printf-style formatting (e.g. __("%s%s") % (a, b)).
                mismatches.append({
                    'file': rpy_file.relative_to(ROOT),
                    'line': line_no,
                    'old': old_text,
                    'new': new_text,
                    'old_ph': old_ph,
                    'new_ph': ['(empty translation)'],
                })
                continue
            new_ph = extract_placeholders(new_text)
            if old_ph != new_ph:
                mismatches.append({
                    'file': rpy_file.relative_to(ROOT),
                    'line': line_no,
                    'old': old_text,
                    'new': new_text,
                    'old_ph': old_ph,
                    'new_ph': new_ph,
                })

    print(f"Checked {checked} translation blocks")
    print(f"Placeholder mismatches: {len(mismatches)}")

    if mismatches:
        export_xlsx(mismatches, "chinese_simplified")
        for m in mismatches[:30]:
            print(f"\n{m['file']}:{m['line']}")
            print(f"  old: {m['old'][:120]!r}")
            print(f"  new: {m['new'][:120]!r}")
            print(f"  old placeholders: {m['old_ph']}")
            print(f"  new placeholders: {m['new_ph']}")
        if len(mismatches) > 30:
            print(f"\n... and {len(mismatches) - 30} more")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
