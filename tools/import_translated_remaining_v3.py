#!/usr/bin/env python3
"""
Import translations from to_translate_remaining_v3.xlsx back into the
corresponding chinese_simplified translation files.

Worksheets:
- strings: game strings from per-file 'translate chinese_simplified strings:' blocks
- common: Ren'Py common strings from common.rpy
- dialogue: dialogue translate blocks
"""

import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Install with: pip install openpyxl")
    exit(1)

NEW_PROJECT = Path(r"C:\Users\akxls\Documents\Code\BK\Brothel_King-pc")
TL_DIR = NEW_PROJECT / "game" / "tl" / "chinese_simplified"
EXCEL_FILE = NEW_PROJECT / "temp" / "translations" / "to_translate_remaining_v3.xlsx"


def parse_quoted_string(s):
    s = s.strip()
    if s.startswith('"""') and s.endswith('"""'):
        return s[3:-3]
    if s.startswith('"') and s.endswith('"'):
        return s[1:-1].replace('\\"', '"').replace('\\n', '\n').replace('\\\\', '\\')
    return s


def format_quoted_string(text):
    if text is None:
        return '""'
    if '\n' in text or '\r' in text:
        safe = text.replace('"""', '\\"""')
        return f'"""{safe}"""'
    else:
        safe = text.replace('"', '\\"').replace('\\', '\\\\')
        return f'"{safe}"'


def import_sheet(ws, sheet_type):
    """Import a sheet of string translations into per-file rpy files."""
    trans = {}
    for row in range(2, ws.max_row + 1):
        file_cell = ws.cell(row, 2).value
        context = ws.cell(row, 3).value
        english = ws.cell(row, 4).value
        chinese = ws.cell(row, 5).value

        if not english or not chinese:
            continue

        key = (file_cell, english)
        trans.setdefault(key, []).append((context, chinese))

    print(f"  Loaded {len(trans)} unique entries from '{sheet_type}' sheet")

    applied = 0
    skipped = 0

    # Group by file
    by_file = {}
    for (file_cell, english), entries in trans.items():
        by_file.setdefault(file_cell, []).append((english, entries))

    for file_cell, entries in by_file.items():
        # Normalize file path: remove leading game/tl/chinese_simplified/ if present
        if file_cell:
            file_cell = file_cell.replace('\\', '/')
            prefix = 'game/tl/chinese_simplified/'
            if file_cell.startswith(prefix):
                file_cell = file_cell[len(prefix):]
        fpath = TL_DIR / file_cell
        if not fpath.exists():
            print(f"    File not found: {fpath}")
            skipped += len(entries)
            continue

        content = fpath.read_text(encoding='utf-8')
        original = content

        # Build english -> chinese dict (use first match)
        eng_to_ch = {}
        for english, ctx_entries in entries:
            eng_to_ch[english] = ctx_entries[0][1]

        block_pattern = re.compile(
            r'^(\s*#\s*[^\n]*\n)?'
            r'(\s*old\s+)("""[\s\S]*?"""|"(?:[^"\\]|\\.)*")\n'
            r'(\s*new\s+)("""[\s\S]*?"""|"(?:[^"\\]|\\.)*")\n',
            re.MULTILINE,
        )

        def replace_block(match):
            old_prefix = match.group(2)
            old_val = match.group(3)
            new_prefix = match.group(4)

            old_text = parse_quoted_string(old_val)
            new_text = parse_quoted_string(match.group(5))

            if old_text in eng_to_ch and new_text == '':
                return f"{old_prefix}{old_val}\n{new_prefix}{format_quoted_string(eng_to_ch[old_text])}\n"
            return match.group(0)

        new_content = block_pattern.sub(replace_block, content)

        if new_content != original:
            fpath.write_text(new_content, encoding='utf-8')

        # Count applied
        for english, chinese in eng_to_ch.items():
            formatted = format_quoted_string(chinese)
            escaped_old = re.escape(english).replace('\n', '\\n')
            if re.search(r'old\s+"' + escaped_old + r'"\s*\n\s*new\s+' + re.escape(formatted), new_content):
                applied += 1
            else:
                skipped += 1

    return applied, skipped


def import_dialogue(ws):
    """Import dialogue sheet into per-file translation files."""
    trans = {}
    for row in range(2, ws.max_row + 1):
        file_cell = ws.cell(row, 2).value
        hash_cell = ws.cell(row, 3).value
        chinese = ws.cell(row, 5).value

        if not file_cell or not hash_cell or not chinese:
            continue

        if hash_cell.startswith('hash='):
            hash_id = hash_cell[5:]
        else:
            hash_id = hash_cell

        trans.setdefault(file_cell, {})[hash_id] = chinese

    print(f"  Loaded {sum(len(v) for v in trans.values())} dialogue entries")

    applied = 0
    for file_cell, hashes in trans.items():
        fpath = TL_DIR / file_cell
        if not fpath.exists():
            print(f"    File not found: {fpath}")
            continue

        content = fpath.read_text(encoding='utf-8')
        for hash_id, chinese in hashes.items():
            pattern = re.compile(
                r'(translate chinese_simplified ' + re.escape(hash_id) + r':\s*\n)'
                r'((?:\s*\n|\s*#.*\n)*)'
                r'(\s*(?:(\w+)\s+)?)("(?:[^"\\]|\\.)*")',
                re.MULTILINE,
            )

            def repl(m):
                return m.group(1) + m.group(2) + m.group(3) + format_quoted_string(chinese)

            new_content = pattern.sub(repl, content)
            if new_content != content:
                content = new_content
                applied += 1

        fpath.write_text(content, encoding='utf-8')

    return applied


def main():
    if not EXCEL_FILE.exists():
        print(f"Excel file not found: {EXCEL_FILE}")
        return

    print(f"Loading {EXCEL_FILE}...")
    wb = openpyxl.load_workbook(EXCEL_FILE)

    print("\nImporting strings...")
    strings_applied, strings_skipped = import_sheet(wb['strings'], 'strings')
    print(f"  Applied: {strings_applied}, Skipped: {strings_skipped}")

    if 'common' in wb.sheetnames:
        print("\nImporting common...")
        common_applied, common_skipped = import_sheet(wb['common'], 'common')
        print(f"  Applied: {common_applied}, Skipped: {common_skipped}")

    if 'dialogue' in wb.sheetnames:
        print("\nImporting dialogue...")
        dialogue_applied = import_dialogue(wb['dialogue'])
        print(f"  Applied: {dialogue_applied}")
    else:
        dialogue_applied = 0
        print("\nNo 'dialogue' sheet found, skipping dialogue import.")

    print("\n" + "=" * 60)
    total = strings_applied + (common_applied if 'common' in wb.sheetnames else 0) + dialogue_applied
    print(f"Total imported: {total}")


if __name__ == "__main__":
    main()
