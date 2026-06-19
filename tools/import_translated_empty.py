#!/usr/bin/env python3
"""
Import translated empty entries from an Excel file back into Ren'Py translation files.

Usage:
    python tools\\import_translated_empty.py [path\\to\\file.xlsx]

If no path is given, defaults to project root\\temp\\translations\\to_translate_empty.xlsx.
"""

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Install with: pip install openpyxl")
    exit(1)

NEW_PROJECT = Path(r"C:\Users\akxls\Documents\Code\BK\Brothel_King-pc")
NEW_TL_DIR = NEW_PROJECT / "game" / "tl" / "chinese_simplified"
DEFAULT_EXCEL = NEW_PROJECT / "temp" / "translations" / "to_translate_empty.xlsx"


def parse_quoted_string(s):
    """Parse a Ren'Py quoted string literal, return raw text."""
    s = s.strip()
    if s.startswith('"""') and s.endswith('"""'):
        return s[3:-3]
    if s.startswith('"') and s.endswith('"'):
        return s[1:-1].replace('\\"', '"').replace('\\n', '\n')
    return s


def format_quoted_string(text):
    """Format raw text as a Ren'Py string literal."""
    if text is None:
        return '""'
    if '\n' in text or '\r' in text:
        safe = text.replace('"""', '\\"""')
        return f'"""{safe}"""'
    else:
        safe = text.replace('"', '\\"')
        return f'"{safe}"'


def load_entries(wb):
    """Load string and dialogue entries from workbook."""
    string_entries = []  # list of (file_path, english, chinese)
    dialogue_entries = []  # list of (file_path, hash_id, chinese)

    if 'strings' in wb.sheetnames:
        ws = wb['strings']
        for row in range(2, ws.max_row + 1):
            file_path = ws.cell(row, 2).value
            english = ws.cell(row, 4).value
            chinese = ws.cell(row, 5).value
            if english and chinese:
                string_entries.append((file_path, english, chinese))

    if 'dialogue' in wb.sheetnames:
        ws = wb['dialogue']
        for row in range(2, ws.max_row + 1):
            file_path = ws.cell(row, 2).value
            hash_cell = ws.cell(row, 3).value
            chinese = ws.cell(row, 5).value
            if hash_cell and chinese:
                if isinstance(hash_cell, str) and hash_cell.startswith('hash='):
                    hash_id = hash_cell[5:]
                else:
                    hash_id = hash_cell
                dialogue_entries.append((file_path, hash_id, chinese))

    return string_entries, dialogue_entries


def import_strings(string_entries):
    """Import strings into the correct .rpy translation files."""
    if not string_entries:
        print("  No string entries to import")
        return 0

    # Group by target file
    by_file = {}
    for file_path, english, chinese in string_entries:
        by_file.setdefault(file_path, []).append((english, chinese))

    total_applied = 0
    for file_path, entries in by_file.items():
        fpath = NEW_TL_DIR / file_path
        if not fpath.exists():
            print(f"  File not found: {fpath}")
            continue

        with open(fpath, 'r', encoding='utf-8') as f:
            content = f.read()

        trans = dict(entries)

        block_pattern = re.compile(
            r'^(\s*#\s*[^\n]*\n)?'
            r'(\s*old\s+)("""[\s\S]*?"""|"(?:[^"\\]|\\.)*")\n'
            r'(\s*new\s+)("""[\s\S]*?"""|"(?:[^"\\]|\\.)*")\n',
            re.MULTILINE,
        )

        def replace_block(match):
            old_prefix = match.group(2)
            old_val = match.group(3)
            new_prefix = match.group(4)
            new_val = match.group(5)

            old_text = parse_quoted_string(old_val)

            if old_text in trans and parse_quoted_string(new_val) == '':
                new_text = trans[old_text]
                return f"{old_prefix}{old_val}\n{new_prefix}{format_quoted_string(new_text)}\n"
            return match.group(0)

        new_content = block_pattern.sub(replace_block, content)

        if new_content != content:
            with open(fpath, 'w', encoding='utf-8') as f:
                f.write(new_content)

        applied = 0
        for old_text, new_text in trans.items():
            formatted_new = format_quoted_string(new_text)
            escaped_old = re.escape(old_text).replace('\n', '\\n')
            if re.search(r'old\s+"' + escaped_old + r'"\s*\n\s*new\s+' + re.escape(formatted_new), new_content):
                applied += 1

        total_applied += applied

    return total_applied


def import_dialogue(dialogue_entries):
    """Import dialogue into per-file translation files."""
    if not dialogue_entries:
        print("  No dialogue entries to import")
        return 0

    # Group by target file
    by_file = {}
    for file_path, hash_id, chinese in dialogue_entries:
        by_file.setdefault(file_path, []).append((hash_id, chinese))

    total_imported = 0
    for file_path, entries in by_file.items():
        fpath = NEW_TL_DIR / file_path
        if not fpath.exists():
            print(f"  File not found: {fpath}")
            continue

        content = fpath.read_text(encoding='utf-8')
        new_content = content

        for hash_id, chinese in entries:
            marker = f'translate chinese_simplified {hash_id}:'
            if marker not in new_content:
                print(f"    Hash not found: {hash_id} in {fpath}")
                continue

            pattern = re.compile(
                r'(translate chinese_simplified ' + re.escape(hash_id) + r':\s*\n)'
                r'((?:\s*\n|\s*#.*\n)*)'
                r'(\s*(?:(\w+)\s+)?)("(?:[^"\\]|\\.)*")',
                re.MULTILINE,
            )

            def repl(m):
                prefix = m.group(1) + m.group(2) + m.group(3)
                return prefix + format_quoted_string(chinese)

            new_content = pattern.sub(repl, new_content)

        if new_content != content:
            fpath.write_text(new_content, encoding='utf-8')

        # Count how many were actually replaced
        imported = 0
        for hash_id, chinese in entries:
            marker = f'translate chinese_simplified {hash_id}:'
            if marker in new_content:
                formatted = format_quoted_string(chinese)
                # Check if the chinese text appears right after the translate block
                if re.search(
                    r'translate chinese_simplified ' + re.escape(hash_id) + r':\s*\n'
                    r'(?:\s*\n|\s*#.*\n)*'
                    r'(?:\s*\w+\s+)?' + re.escape(formatted),
                    new_content,
                ):
                    imported += 1

        total_imported += imported

    return total_imported


def main():
    excel_file = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_EXCEL
    if not excel_file.exists():
        print(f"Excel file not found: {excel_file}")
        return

    print(f"Loading {excel_file}...")
    wb = openpyxl.load_workbook(excel_file)

    string_entries, dialogue_entries = load_entries(wb)

    print(f"\nImporting strings... ({len(string_entries)} entries)")
    strings_imported = import_strings(string_entries)
    print(f"  Applied: {strings_imported}")

    print(f"\nImporting dialogue... ({len(dialogue_entries)} entries)")
    dialogue_imported = import_dialogue(dialogue_entries)
    print(f"  Applied: {dialogue_imported}")

    print("\n" + "=" * 60)
    print(f"Total applied: {strings_imported + dialogue_imported}")


if __name__ == "__main__":
    main()
