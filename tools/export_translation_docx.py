import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).parent.parent
XLSX_PATH = ROOT / 'to_translate_remaining.xlsx'
DOCX_PATH = ROOT / 'to_translate_remaining.docx'


def main():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX_PATH)
    
    doc = Document()
    
    # Title
    title = doc.add_heading('Brothel King - 待翻译剧情对话', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('说明：请在「中文翻译」列填写对应的中文翻译。翻译完成后保存此 docx 文件即可。')
    doc.add_paragraph()
    
    sheet_order = ['chapter1', 'chapter2', 'chapter3', 'story', 'city', 'day',
                   'interact', 'interact_free', 'help', 'intro', 'main',
                   'powers', 'security', 'start', 'kite1', 'kite2']
    
    sheet_names_cn = {
        'chapter1': '第一章',
        'chapter2': '第二章',
        'chapter3': '第三章',
        'story': '故事事件',
        'city': '城市事件',
        'day': '日常事件',
        'interact': '互动对话',
        'interact_free': '自由互动',
        'help': '帮助文本',
        'intro': '序章',
        'main': '主线',
        'powers': '能力',
        'security': '安全事件',
        'start': '开始',
        'kite1': '风筝女孩1',
        'kite2': '风筝女孩2',
    }
    
    for sheet_name in sheet_order:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))[1:]  # skip header
        
        if not rows:
            continue
        
        # Section heading
        cn_name = sheet_names_cn.get(sheet_name, sheet_name)
        doc.add_heading(f'{cn_name} ({sheet_name}) - {len(rows)} 条', level=1)
        
        # Table: English | Chinese | Speaker | Label
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = '英文原文'
        hdr_cells[1].text = '中文翻译'
        hdr_cells[2].text = '角色'
        hdr_cells[3].text = '标签'
        
        # Make header bold
        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        
        for row_data in rows:
            if not row_data:
                continue
            en = str(row_data[0]) if row_data[0] else ''
            speaker = str(row_data[3]) if len(row_data) > 3 and row_data[3] else ''
            label = str(row_data[4]) if len(row_data) > 4 and row_data[4] else ''
            
            if not en.strip():
                continue
            
            row_cells = table.add_row().cells
            row_cells[0].text = en
            row_cells[1].text = ''  # Empty for translation
            row_cells[2].text = speaker
            row_cells[3].text = label
        
        doc.add_paragraph()  # spacing between sections
    
    doc.save(DOCX_PATH)
    print(f'Saved {DOCX_PATH}')
    
    # Count total
    total = sum(len(list(wb[s].iter_rows(values_only=True))[1:]) for s in sheet_order if s in wb.sheetnames)
    print(f'Total entries exported: {total}')


if __name__ == '__main__':
    main()
