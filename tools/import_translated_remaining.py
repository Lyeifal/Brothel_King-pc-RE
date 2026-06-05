import re
import os
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).parent.parent
STRINGS_PATH = ROOT / 'game/tl/chinese_simplified/strings.rpy'
TL_DIR = ROOT / 'game/tl/chinese_simplified'
XLSX_ORIG = ROOT / 'to_translate_remaining.xlsx'
XLSX_TRANS = ROOT / 'to_translated_remaining.xlsx'


def quote_for_rpy(text):
    """Escape text for Ren'Py string."""
    text = text.replace('\\', '\\\\')
    text = text.replace('"', '\\"')
    text = text.replace('\n', '\\n')
    return text


def unquote_rpy(text):
    """Unescape Ren'Py string."""
    text = text.replace('\\n', '\n')
    text = text.replace('\\"', '"')
    text = text.replace('\\\\', '\\')
    return text


def extract_quoted(text):
    """Extract content from first quoted string using proper escape handling."""
    match = re.search(r'"((?:\\.|[^"\\])*)"', text)
    if match:
        return match.group(1)
    return ''


def import_strings(orig_ws, trans_ws):
    """Import strings sheet into strings.rpy using row alignment."""
    updates = {}
    skipped = 0
    for row_idx in range(2, orig_ws.max_row + 1):
        orig_val = orig_ws.cell(row=row_idx, column=1).value
        trans_val = trans_ws.cell(row=row_idx, column=2).value
        
        if orig_val is None:
            continue
        orig_str = str(orig_val).strip()
        trans_str = str(trans_val).strip() if trans_val else ''
        
        if not trans_str or trans_str == orig_str:
            skipped += 1
            continue
        updates[orig_str] = trans_str
    
    print(f"  Loaded {len(updates)} string updates (skipped {skipped} empty/unchanged)")
    if not updates:
        return 0
    
    with open(STRINGS_PATH, 'r', encoding='utf-8') as f:
        content = f.read()
    
    pattern = re.compile(
        r'(\s*#.*?)\n\s*old\s+"(.*?)"\s*\n\s*new\s+"(.*?)"\s*\n',
        re.MULTILINE | re.DOTALL
    )
    
    match_count = 0
    def replacer(m):
        nonlocal match_count
        comment = m.group(1).strip()
        old_quoted = m.group(2)
        old_text = unquote_rpy(old_quoted)
        if old_text in updates:
            new_text = updates[old_text]
            new_quoted = quote_for_rpy(new_text)
            comment = comment.replace('TODO', '').strip()
            match_count += 1
            return f'    {comment}\n    old "{old_quoted}"\n    new "{new_quoted}"\n'
        return m.group(0)
    
    new_content = pattern.sub(replacer, content)
    
    if new_content != content:
        with open(STRINGS_PATH, 'w', encoding='utf-8') as f:
            f.write(new_content)
        print(f"  Updated strings.rpy ({match_count} entries matched and updated)")
    else:
        print("  No changes made to strings.rpy")
    
    return match_count


def import_dialogue(orig_ws, trans_ws, rpy_path):
    """Import a dialogue sheet using row alignment by label."""
    rpy_path = ROOT / rpy_path
    updates = {}
    skipped = 0
    
    for row_idx in range(2, orig_ws.max_row + 1):
        orig_val = orig_ws.cell(row=row_idx, column=1).value
        trans_val = trans_ws.cell(row=row_idx, column=2).value
        label = orig_ws.cell(row=row_idx, column=5).value
        
        if orig_val is None or label is None:
            continue
        orig_str = str(orig_val).strip()
        trans_str = str(trans_val).strip() if trans_val else ''
        label_str = str(label).strip()
        
        if not trans_str or trans_str == orig_str:
            skipped += 1
            continue
        updates[label_str] = (orig_str, trans_str)
    
    print(f"  Loaded {len(updates)} dialogue updates (skipped {skipped} empty/unchanged)")
    if not updates:
        return 0
    
    with open(rpy_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    match_count = 0
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        match = re.match(r'translate\s+chinese_simplified\s+(\w+)\s*:\s*$', stripped)
        if match:
            label = match.group(1)
            if label in updates:
                orig_text, trans_text = updates[label]
                # Find the first non-comment, non-empty line after the label
                j = i + 1
                while j < len(lines):
                    s = lines[j].strip()
                    if s and not s.startswith('#'):
                        # This is the dialogue line
                        quoted = extract_quoted(s)
                        if quoted and unquote_rpy(quoted).strip() == orig_text.strip():
                            # Replace the quoted text
                            new_quoted = quote_for_rpy(trans_text)
                            lines[j] = lines[j].replace(f'"{quoted}"', f'"{new_quoted}"', 1)
                            match_count += 1
                        break
                    j += 1
        i += 1
    
    if match_count > 0:
        with open(rpy_path, 'w', encoding='utf-8') as f:
            f.writelines(lines)
        print(f"  Updated {rpy_path} ({match_count} entries matched and updated)")
    else:
        print(f"  No changes made to {rpy_path}")
    
    return match_count


def main():
    if len(sys.argv) > 1:
        trans_file = ROOT / sys.argv[1]
    else:
        trans_file = XLSX_TRANS
    
    print(f"Loading original: {XLSX_ORIG}")
    print(f"Loading translated: {trans_file}")
    orig_wb = openpyxl.load_workbook(XLSX_ORIG)
    trans_wb = openpyxl.load_workbook(trans_file)
    
    file_map = {
        'chapter1': str(ROOT / 'game/tl/chinese_simplified/BKchapter1.rpy'),
        'chapter2': str(ROOT / 'game/tl/chinese_simplified/BKchapter2.rpy'),
        'chapter3': str(ROOT / 'game/tl/chinese_simplified/BKchapter3.rpy'),
        'story': str(ROOT / 'game/tl/chinese_simplified/BKstory_events.rpy'),
        'city': str(ROOT / 'game/tl/chinese_simplified/BKcity_events.rpy'),
        'day': str(ROOT / 'game/tl/chinese_simplified/BKday_events.rpy'),
        'interact': str(ROOT / 'game/tl/chinese_simplified/BKinteractions.rpy'),
        'interact_free': str(ROOT / 'game/tl/chinese_simplified/BKinteractions_free.rpy'),
        'help': str(ROOT / 'game/tl/chinese_simplified/BKhelp.rpy'),
        'intro': str(ROOT / 'game/tl/chinese_simplified/BKintro.rpy'),
        'main': str(ROOT / 'game/tl/chinese_simplified/BKmain.rpy'),
        'powers': str(ROOT / 'game/tl/chinese_simplified/BKpowers.rpy'),
        'security': str(ROOT / 'game/tl/chinese_simplified/BKsecurity.rpy'),
        'start': str(ROOT / 'game/tl/chinese_simplified/BKstart.rpy'),
        'kite1': str(ROOT / 'game/tl/chinese_simplified/kite_jobgirl 1_riddle.rpy'),
        'kite2': str(ROOT / 'game/tl/chinese_simplified/kite_jobgirl 2_beach.rpy'),
    }
    
    total = 0
    
    # Import strings
    print("\nImporting string translations into strings.rpy...")
    orig_ws = orig_wb['strings']
    trans_ws = trans_wb[trans_wb.sheetnames[0]]  # First sheet = strings
    total += import_strings(orig_ws, trans_ws)
    
    # Import dialogues
    for i, sheet_name in enumerate(orig_wb.sheetnames[1:], start=1):
        if sheet_name not in file_map:
            continue
        rpy_path = file_map[sheet_name]
        trans_sheet_name = trans_wb.sheetnames[i]
        print(f"\nImporting {sheet_name} translations into {rpy_path}...")
        total += import_dialogue(orig_wb[sheet_name], trans_wb[trans_sheet_name], rpy_path)
    
    print(f"\n{'='*60}")
    print(f"Total translations imported: {total}")
    print(f"{'='*60}")
    
    # Delete .rpyc cache
    print("\nDeleting .rpyc cache files...")
    tl_dir = TL_DIR
    for rpyc in tl_dir.glob('*.rpyc'):
        rpyc.unlink()
        print(f"  Removed {rpyc}")
    
    print("\nDone! Launch the game to verify translations.")


if __name__ == '__main__':
    main()
